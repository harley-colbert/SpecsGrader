from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from backend.app.config.xlsx_contract import (
    DEPT_COL,
    RISK_LEVEL_COL,
    SPEC_TEXT_COL,
    SPECIFIC_RISK_COL,
    is_medium_plus,
    normalize_risk_level,
)
from backend.app.services.specific_risk_service import generate_specific_risk


def _select_sheet(sheet_names: Iterable[str]) -> str:
    for sheet in sheet_names:
        if "standards risk matrix" in sheet.lower():
            return sheet
    raise ValueError("No sheet containing 'Standards Risk Matrix' found")


def write_classify_predictions_to_xlsx(
    source_path: str,
    results: list[dict[str, object]],
    overwrite_predictions: bool = True,
    overwrite_specific_risk: bool = True,
) -> Path:
    path = Path(source_path)
    if path.suffix.lower() not in {".xlsx", ".xls"}:
        raise ValueError("Classify export supports XLSX only")

    wb = load_workbook(path)
    sheet = wb[_select_sheet(wb.sheetnames)]
    by_row = {int(result["source_row"]): result for result in results if result.get("source_row")}

    for row_num, result in by_row.items():
        spec_cell = f"{SPEC_TEXT_COL}{row_num}"
        spec_value = sheet[spec_cell].value
        if spec_value is None or str(spec_value).strip() == "":
            continue

        level_cell = f"{RISK_LEVEL_COL}{row_num}"
        dept_cell = f"{DEPT_COL}{row_num}"
        specific_risk_cell = f"{SPECIFIC_RISK_COL}{row_num}"
        if not overwrite_predictions:
            if sheet[level_cell].value or sheet[dept_cell].value:
                continue

        pred_level = result.get("pred_level") or ""
        pred_dept = result.get("pred_dept") or ""
        sheet[level_cell].value = pred_level
        sheet[dept_cell].value = pred_dept

        normalized_level = normalize_risk_level(pred_level) or ""
        if not is_medium_plus(normalized_level):
            sheet[specific_risk_cell].value = ""
            continue

        if not overwrite_specific_risk:
            existing_specific = sheet[specific_risk_cell].value
            if existing_specific is not None and str(existing_specific).strip():
                continue

        sheet[specific_risk_cell].value = generate_specific_risk(
            str(spec_value),
            normalized_level,
            str(pred_dept),
        )

    output_dir = path.parent / "exports"
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    output_path = output_dir / f"{path.stem}_classified_{timestamp}{path.suffix}"
    wb.save(output_path)
    return output_path
